"""Builds examples/sample_grid_timeline.xlsx: a week-by-week team schedule (one column per week).

Mirrors a common production-schedule layout: week dates across the top (no year), a merged
phase banner, a milestone row, a "Week N" row, then one row per person with their task each week.
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

weeks = ["August 23", "May 4", "May 11", "August 30", "May 24", "September 6", "September 13",
         "September 20", "September 27", "October 4", "October 11", "October 18", "October 25",
         "November 1st", datetime(2026, 11, 8), "November 15", "November 22", "November 29",
         "December 6", "December 13"]
hidden = {"May 4", "May 11", "May 24"}  # old columns left hidden in the sheet

wb = Workbook()
ws = wb.active
ws.title = "Schedule"
ws["A1"] = "GAME STUDIO"
ws.merge_cells("A1:B3")
col_of = {}
for i, label in enumerate(weeks, start=3):
    ws.cell(row=1, column=i, value=label)
    col_of[label if isinstance(label, str) else "November 8"] = i
    if label in hidden:
        ws.column_dimensions[get_column_letter(i)].hidden = True
    ws.cell(row=4, column=i, value=f"Week {i - 2}")
c = lambda name: get_column_letter(col_of[name])

ws[f"{c('August 23')}2"] = "PRE PRODUCTION"
ws.merge_cells(f"{c('August 23')}2:{c('October 4')}2")
ws[f"{c('October 11')}2"] = "ALPHA PHASE"
ws.merge_cells(f"{c('October 11')}2:{c('December 6')}2")
ws[f"{c('December 13')}2"] = "Beta"
for week, text in {"September 20": "Pitch Day!", "October 4": "Design Pillars", "October 11": "Dev Begins",
                   "November 1st": "Playtest", "November 15": "Playtest", "December 6": "Alpha Ends"}.items():
    ws[f"{c(week)}3"] = text
ws["A4"], ws["B4"] = "Track / Focus", "Name"

people = [
    ("Art", "Ava", {"October 4": "Desert Level 3D Boss Concept", "October 18": "Boss Textures"}),
    ("Art", "Ben", {"October 4": "More VFX Foam", "October 18": "Foam Shader Polish"}),
    ("Tech Art", "Cam", {"October 4": "Desert Level Region Mockup", "October 25": "Rigging Pass"}),
    ("Design", "Dee", {"October 4": "Desert Levels Blueprint Iteration", "October 18": "Playtest Plan"}),
    ("Engineering", "Eli", {"October 4": "Desert Final Level Boss Fight", "October 18": "Enemy AI Pass"}),
    ("Production", "Fay", {"October 4": "Presentation, Timeline, and Task updates", "October 18": "Sprint Review"}),
]
for r, (track, name, tasks) in enumerate(people, start=5):
    ws.cell(row=r, column=1, value=track)
    ws.cell(row=r, column=2, value=name)
    ws.cell(row=r, column=col_of["August 23"], value="Ideation & Gen. Plan")
    for week, text in tasks.items():
        ws.cell(row=r, column=col_of[week], value=text)
ws[f"{c('October 11')}5"] = "Fall Break!"
ws.merge_cells(f"{c('October 11')}5:{c('October 11')}{4 + len(people)}")

out = Path(__file__).with_name("sample_grid_timeline.xlsx")
wb.save(out)
print(out)
